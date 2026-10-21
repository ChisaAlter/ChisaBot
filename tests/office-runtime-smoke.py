"""Run inside the built image as pwuser: real Excel/Word round trips, independent totals."""
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook, load_workbook
from docx import Document

with TemporaryDirectory(prefix="office-runtime-") as folder:
    root = Path(folder)
    xlsx = root / "预算.xlsx"
    word = root / "复盘.docx"
    rows = [("研发", 12000, 13500), ("设计", 5000, 4200), ("测试", 3000, 2800), ("运营", 4000, 3900), ("培训", 2000, 1800)]
    wb = Workbook()
    sheet = wb.active
    sheet.title = "预算"
    sheet.append(["类别", "预算", "实际", "差额"])
    for index, row in enumerate(rows, 2):
        sheet.append([*row, f"=C{index}-B{index}"])
    sheet.append(["总计", "=SUM(B2:B6)", "=SUM(C2:C6)", "=SUM(D2:D6)"])
    info = wb.create_sheet("说明")
    info.append(["项目", "海盐"])
    info.append(["负责人", "林晓"])
    wb.save(xlsx)
    reopened = load_workbook(xlsx, data_only=False)
    read_rows = list(reopened["预算"].iter_rows(min_row=2, max_row=6, values_only=True))
    totals = sum(row[1] for row in read_rows), sum(row[2] for row in read_rows)
    assert totals == (26000, 26200)
    assert sum(row[2] - row[1] for row in read_rows) == 200
    assert reopened["预算"]["D7"].value == "=SUM(D2:D6)"
    assert all(row[3] == f"=C{index}-B{index}" for index, row in enumerate(read_rows, 2))
    document = Document()
    document.add_heading("海盐项目复盘", 0)
    document.add_paragraph("负责人：林晓")
    document.add_paragraph(f"总预算 {totals[0]}，总实际 {totals[1]}，差额 {totals[1] - totals[0]}。")
    document.save(word)
    text = "\n".join(paragraph.text for paragraph in Document(word).paragraphs)
    assert "26200" in text and "林晓" in text and "差额 200" in text
    print("OFFICE_RUNTIME_PASS: Excel formulas and numeric rows retained; Word agrees: 26000 / 26200 / 200")
